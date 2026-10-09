"""Memória da ligação com talas: quadro completo e contas governantes."""
from io import BytesIO
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from .models import STEELS,KGF,VERSION
from .double_cover import NBR,ASSEMBLY_NOTE
from .double_cover_drawing import image_bytes
from .report import table,number,display,decimal_text,hyperlink,BRAND,LINK,_math,group,sub,frac,power,radical

REFERENCES=[
    NBR+' — itens 5.3, 6.1.5, 6.2, 6.3 e 6.5: materiais, parafusos, soldas e elementos de ligação.',
    'AISC 360-22, F11: resistência à flexão de barras chatas. A resistência nominal complementar é dividida por γa1; não se utiliza conversão global de tabelas LRFD.',
    'AISC Design Examples v16.0, P901-23W, Parte II, exemplos de emendas de alma: equilíbrio dos grupos simétricos e estados limite dos componentes.',
    'Fortney e Thornton (2016), Analysis and Design of Stabilizer Plates in Single-Plate Shear Connections, Engineering Journal 53(1), 1–28. DOI: 10.62913/engj.v53i1.1095. Fundamenta a atenção ao momento local e à estabilização; não valida esta geometria com talas.',
    'Sapkota et al. (2024), Behavior of Extended Single-Plate Shear Connections Subjected to Combined Shear and Compression Forces Using FEA, Engineering Journal 61(4), 193–216. DOI: 10.62913/engj.v61i4.1332. Referência de comportamento sob compressão, sem transposição direta de capacidades.'
]

def equation(doc,row):
    i=row.id;g1=sub('γ','a1');g2=sub('γ','a2');fy=sub('f','y');fu=sub('f','u');Ag=sub('A','g');An=sub('A','n')
    if i.startswith('dc_bolt_'):
        expr=group(sub('F','Rd'),' = ',frac(group('2',sub('k','pega'),sub('k','calço'),'α',sub('A','b'),sub('f','ub')),g2))
    elif i.startswith('dc_bearing_'):
        expr=group(sub('F','Rd'),' = ',frac(group(sub('k','calço'),sub('n','t'),' min(1,2',sub('l','c'),'t',fu,'; 2,4',sub('d','b'),'t',fu,')'),g2))
    elif i.startswith('dc_block_'):
        expr=group('η = ',frac('Σ|Fx,i|',sub('R','N,bloco')),' + ',frac('Σ|Fy,i|',sub('R','V,bloco')),' ≤ 1')
    elif i.startswith('dc_yield_'):
        expr=group(sub('σ','eq'),' = ',radical(group(power(group('(',frac('|N|',Ag),' + ',frac('|M|',sub('W','g')),')')),' + 3',power(group('(',frac('1,5|V|',Ag),')')))),' ≤ ',frac(fy,g1))
    elif i.startswith('dc_net_'):
        expr=group('η = ',frac(group('max(0; ',frac('N',sub('A','e')),' + ',frac('|M|',sub('W','n')),')'),frac(fu,g2)),' + ',frac('|V|',frac(group('0,6',fu,An),g2)),' ≤ 1')
    elif i.startswith('dc_stability_'):
        expr=group('η = ',frac('|N|',sub('N','Rd')),' + ',frac('|M|',sub('M','Rd')),' + ',frac('|V|',sub('V','Rd')),' ≤ 1')
    elif i=='dc_stiffener':
        expr=group('η = ',frac('|H|s/4',sub('M','Rd')),' + ',frac('|H|/2',sub('V','Rd')),' ≤ 1')
    elif i.startswith('dc_weld_'):
        expr=group(sub('F','Rd'),' = ',frac(group('2',frac('w',radical('2')),'L·0,6',sub('f','w')),sub('γ','w2')))
    elif i.startswith('dc_base_'):
        expr=group(sub('F','Rd'),' = L min(',frac(group('0,6',fy,'t'),g1),'; ',frac(group('0,6',fu,'t'),g2),')')
    else:
        expr=group(sub('V','Rd'),' = 2H',sub('t','w'),' min(',frac(group('0,6',fy),g1),'; ',frac(group('0,6',fu),g2),')')
    _math(doc,expr)
    if i.startswith('dc_block_'):
        _math(doc,group('R = ',frac(group(sub('n','t'),' min(0,6',fu,sub('A','nv'),' + 0,5',fu,sub('A','nt'),'; 0,6',fy,sub('A','gv'),' + 0,5',fu,sub('A','nt'),')'),g2)))

def create_report(c,r,detailed=False):
    if not r.checks or any(x.severity=='error' for x in r.issues):raise ValueError('A geometria deve ser válida e possuir ações para gerar a memória.')
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
        style=doc.styles[name];style.font.name='Arial';style.font.color.rgb=RGBColor(0,0,0)
    for style in doc.styles:
        for border in style.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    normal=doc.styles['Normal'];normal.font.size=Pt(9.5);normal.paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(20);doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 2'].font.size=Pt(10)
    doc.core_properties.title='Ligação de viga à alma de coluna com duas talas'
    doc.core_properties.author='LRO Soluções de engenharia LTDA.'
    doc.add_paragraph('Ligação de viga à alma de coluna',style='Title')
    doc.add_paragraph('Duas talas parafusadas e nervura soldada',style='Subtitle')
    doc.add_paragraph(c.project)
    doc.add_paragraph(c.beam.name+' → '+c.support.name)
    p=doc.add_paragraph();p.add_run(r.status).bold=True
    p.add_run(' | Maior índice calculado: '+number(r.governing.ratio,3))
    doc.add_paragraph(ASSEMBLY_NOTE+' Os índices abaixo se referem aos componentes do modelo local e não aprovam esses mecanismos pendentes.')
    doc.add_picture(BytesIO(image_bytes(c)),width=Inches(6.95))
    doc.add_paragraph('As talas são independentes do conjunto soldado: cada grupo de parafusos transmite V e N integrais; cada tala recebe metade. A viga não encosta na coluna ou na nervura. Encontro ortogonal, com as peças centradas na altura da viga.')
    doc.add_paragraph('O desenho acompanha a geometria informada. É ilustrativo e não substitui o detalhamento de fabricação, a definição dos comprimentos dos parafusos ou a especificação de montagem.')
    doc.add_page_break()
    doc.add_paragraph('Dimensões e materiais',style='Heading 1')
    rows=[
        ('Viga apoiada',f'd={c.beam.d:g}; bf={c.beam.bf:g}; tw={c.beam.tw:g}; tf={c.beam.tf:g}',c.beam_steel),
        ('Coluna',f'd={c.support.d:g}; bf={c.support.bf:g}; tw={c.support.tw:g}; tf={c.support.tf:g}',c.support_steel),
        ('Nervura vertical',f'H={c.root_height:g}; h={c.hp:g}; L={number(c.fin_length)}; t={number(c.fin_t,4)}',c.fin_steel),
        ('Duas talas iguais',f'{number(c.cover_length)} × {number(c.hp)} × {number(c.cover_t,4)}',c.cover_steel),
        ('Quatro horizontais',f'b={number(c.projection)}; s={number(c.stiffener_span)}; t={number(c.stiffener_t,4)}; c={c.corner_clip:g}',c.stiffener_steel),
        ('Parafusos',f'2 grupos × {c.n} unidades; Ø={number(c.db,3)}; 2 planos de corte',c.bolt),
        ('Furos padrão',f'Ø={number(c.dh,4)}; dedução líquida={number(c.dh_net,4)}','Broca' if c.drilled else 'Acréscimo de 2 mm'),
        ('Posições',f'g={c.gap:g}; u={c.extension:g}; p={c.pitch:g}; e={c.edge_x:g}; eᵥ={c.edge_v:g}','mm'),
        ('Soldas duplas',f'w nervura/alma={c.weld_web:g}; nervura/horizontal={c.weld_fin_stiffener:g}; horizontal/coluna={c.weld_stiffener_column:g}',f'fw={c.fw:g} MPa'),
        ('Calços simétricos',f'2 × {number(c.shim_each,3)} no lado da {c.shim_side}' if c.shim_each else f'Ajuste de montagem; diferença={number(c.thickness_difference,3)}','Sem crédito resistente'),
    ]
    rows=[(a,decimal_text(b),d) for a,b,d in rows]
    table(doc,['Componente','Dimensões em mm','Material / critério'],rows,[1.25,3.55,2.15])
    steels=list(dict.fromkeys([c.beam_steel,c.support_steel,c.fin_steel,c.cover_steel,c.stiffener_steel]))
    doc.add_paragraph('Resistências características: '+'; '.join(f'{s}: fy={STEELS[s].fy:g} MPa, fu={STEELS[s].fu:g} MPa' for s in steels)+'. E=200.000 MPa; γa1=1,10; γa2=γw2=1,35.')
    doc.add_paragraph('Ações e equilíbrio',style='Heading 1')
    rows=[]
    for key,label in [('actual_actions','Entrada'),('actions','Conferência governante')]:
        a=r.geometry[key]
        rows.append((label,number(a['V']/KGF),number(a['N']/KGF),number(a['M']/(KGF*1000)),number(a['top']/KGF),number(a['bottom']/KGF)))
    table(doc,['Caso','V kgf','N kgf','M local kgf·m','H superior kgf','H inferior kgf'],rows,[1.45,.9,.9,1.3,1.2,1.2])
    doc.add_paragraph('N positivo indica tração; negativo, compressão. Entradas já majoradas. A conferência de resistência mínima de 45 kN preserva a direção N/V, com fator '+number(r.minimum_factor,3)+'. O momento local deriva da excentricidade; não há momento externo de entrada.')
    _math(doc,group(sub('e','g'),' = e + g/2;   ',sub('M','g'),' = V',sub('e','g'),';   ',sub('M','col'),' = V(b + u + g/2)'))
    _math(doc,group(sub('H','sup'),' = N/2 + ',frac(sub('M','col'),'H'),';   ',sub('H','inf'),' = N/2 − ',frac(sub('M','col'),'H')))
    doc.add_paragraph('b é a projeção da mesa medida da face da alma da coluna; H é a distância livre entre horizontais adotada como braço do binário; e é a borda horizontal; g é a folga nervura–viga. O momento global no eixo da coluna inclui também a distância da face ao eixo da alma. Esses efeitos devem ser compatibilizados com o modelo da estrutura.')
    doc.add_page_break()
    doc.add_paragraph('Quadro de verificações dos componentes',style='Heading 1')
    doc.add_paragraph('Índice η = solicitação/resistência, ou valor da interação. Limite: 1,000. Resultados da conferência governante; “Atende” é exclusivamente o resultado da linha.')
    rows=[]
    for row in r.checks:
        sd,un=display(row.demand,row.unit);rd,_=display(row.resistance,row.unit)
        rows.append((row.name,sd,rd,un,number(row.ratio,3),'Sim' if row.passed else 'Não'))
    table(doc,['Verificação','Solicitação','Resistência','Unid.','η','Atende'],rows,[3.25,.85,.85,.6,.65,.75])
    doc.add_paragraph('Caso governante: '+r.governing.case+'. As verificações da entrada original também são executadas. Como o modelo é homogêneo nas ações, a conferência mínima governa quando seu fator supera 1.')
    doc.add_page_break()
    doc.add_paragraph('Equações e substituições',style='Heading 1')
    doc.add_paragraph('Unidades nas substituições: N, mm e MPa. A tabela anterior converte forças para kgf. '+('Todas as verificações são desenvolvidas.' if detailed else 'Mostra-se o item de maior índice em cada família; o quadro anterior contém todas as verificações.'))
    selected=[]
    for prefix in ['dc_bolt_','dc_bearing_','dc_block_','dc_yield_','dc_net_','dc_stability_','dc_stiffener','dc_weld_','dc_base_','dc_col_web']:
        family=[x for x in r.checks if x.id.startswith(prefix)]
        selected.extend(family if detailed else [max(family,key=lambda x:x.ratio)])
    for row in selected:
        start=len(doc.paragraphs)
        doc.add_paragraph(row.name+' · η='+number(row.ratio,3),style='Heading 2')
        equation(doc,row)
        doc.add_paragraph(decimal_text(row.substitution))
        doc.add_paragraph(row.variables)
        p=doc.add_paragraph(row.reference);p.paragraph_format.space_after=Pt(8)
        for run in p.runs:run.font.size=Pt(8)
        for p in doc.paragraphs[start:-1]:p.paragraph_format.keep_with_next=True
    doc.add_paragraph('Definição das variáveis e hipóteses',style='Heading 1')
    doc.add_paragraph('Ag, An e Ae: áreas bruta, líquida e efetiva; Wg e Wn: módulos elásticos bruto e líquido; Ab: área nominal do parafuso; fy, fu e fub: resistências do aço e do parafuso; db: diâmetro; lc: ligamento livre; nt: quantidade de chapas; α=0,45 com rosca e 0,56 sem rosca; kpega e kcalço: reduções por pega e enchimento; w: perna; L: comprimento útil da solda; χ: fator de instabilidade; K: fator de comprimento; NRd, MRd e VRd: resistências de cálculo.')
    doc.add_paragraph('O grupo elástico usa Fx,i=N/n−Mg·yi/Σyi² e Fy,i=V/n. Para o bloco U axial: Agv=2et; Anv=2(e−dn/2)t; Ant=(n−1)(p−dn)t. Para o bloco L vertical: lv=ev+(n−1)p; Agv=lv·t; Anv=[lv−(n−0,5)dn]t; Ant=(e−dn/2)t. As áreas gv/nv/nt são, respectivamente, bruta ao corte, líquida ao corte e líquida à tração; dn é o diâmetro de dedução líquida. A soma de duas peças iguais é feita por nt=2.')
    doc.add_paragraph('Seções: Ag=nt·h·t; An=nt·(h−n·dn)t; Wg=nt·t·h²/6. Nas talas Ae=min(An;0,85Ag). Wn=nt·[t·h³/12−Σt(dn³/12+dn·yi²)]/(h/2). Nos trechos sem furos, n=0. As interações lineares e as faixas resistentes são hipóteses deste modelo, não uma equação normativa única para o nó completo.')
    doc.add_paragraph('Compressão: Ne=π²EI/(KL)²; λ0=√(Agfy/Ne); χ=0,658^(λ0²) se λ0≤1,5 e 0,877/λ0² nos demais casos. Adota-se a curva inclusive para barras curtas. As chapas são idealizadas como barras chatas isoladas, Q=1, sem ganho por solidarização das talas. K=1 entre grupos das talas; K=2 nas faixas da nervura e da alma da viga. Esses comprimentos são conservadores para os componentes idealizados, mas não demonstram o modo acoplado do conjunto.')
    doc.add_paragraph('Na flexão, o complemento AISC F11 usa λ=Lh/t² e Cb=1: Mn=Mp até 0,08E/fy; no trecho intermediário, Mn=min[Mp;(1,52−0,274λfy/E)My]; acima de 1,9E/fy, Mn=min[Mp;1,9E·Wg/λ]. MrD=Mn/γa1. A faixa horizontal é biapoiada entre mesas, com carga no centro e largura útil b−2c. Não se atribui ganho resistente ao par de horizontais oposto.')
    for issue in r.issues:
        if issue.severity in ('pending','error'):doc.add_paragraph('Pendência: '+issue.text)
    doc.add_paragraph('Premissas: viga com contenção eficaz; juntas internas mesa–alma dos perfis soldados com penetração total e metal de adição compatível; ligação por contato, soldagem em oficina. A resistência/rigidez de estabilização dos horizontais e suas soldas para forças fora do plano não está abrangida pela verificação de H no plano. Análise global da coluna, atrito, fadiga, vibração e carregamento cíclico fora do escopo.')
    if c.notes:doc.add_paragraph('Observações do projeto: '+c.notes)
    doc.add_paragraph('Referências',style='Heading 1')
    for ref in REFERENCES:doc.add_paragraph(ref)
    doc.add_paragraph(BRAND)
    p=doc.add_paragraph('LRO Ligações '+VERSION+' · ');hyperlink(p,'LinkedIn — Lucas Oliveira',LINK)
    out=BytesIO();doc.save(out);return out.getvalue()
