"""Regressão do projeto real e das opções de montagem, sem novas resistências."""
from pathlib import Path
from dataclasses import replace
from io import BytesIO
import json
import pytest
from docx import Document
from streamlit.testing.v1 import AppTest
from lro.models import Connection, COMPLEMENT_ASSUMPTION, INTERFERENCE_ASSUMPTION
from lro.engine import evaluate
from lro.examples import presets
from lro.report import create_report

ROOT=Path(__file__).resolve().parents[1]
PROJECT=ROOT/'examples/Projeto_usuario_W150_W250.json'

def user_case():
    return Connection.from_dict(json.loads(PROJECT.read_text()))

def test_actual_user_project_is_calculated_without_hiding_real_failures():
    c=user_case();r=evaluate(c)
    assert len(r.checks)==24 and len(r.actual)==24
    assert r.status=='NÃO ATENDE'
    assert not any(i.severity=='error' for i in r.issues)
    assert sum(i.severity=='interference' for i in r.issues)==2
    assert r.governing.id=='bearing_beam' and r.governing.ratio>1
    assert c.V==pytest.approx(8139.5195) and c.N==pytest.approx(18632.635)
    assert r.minimum_factor==pytest.approx(45000/(c.V*c.V+c.N*c.N)**.5)
    d=Document(BytesIO(create_report(c,r)))
    text='\n'.join(p.text for p in d.paragraphs)
    assert 'NÃO ATENDE' in text and 'Avisos de montagem' in text
    assert text.index('Avisos de montagem')>text.index('Premissas adotadas')
    assert INTERFERENCE_ASSUMPTION in text

@pytest.mark.parametrize('width,clip,weld',[(45.,20.,5.),(60.,5.,2.),(72.,20.,10.)])
def test_complement_does_not_add_checks_or_increase_any_capacity(width,clip,weld):
    c=user_case();base=evaluate(c)
    full=replace(c,plate_shape='between_flanges',root_width=width,corner_clip=clip,flange_weld=weld)
    r=evaluate(full)
    assert r.checks==base.checks and r.actual==base.actual
    assert any(i.text==COMPLEMENT_ASSUMPTION for i in r.issues)
    assert Connection.from_dict(full.to_dict())==full
    assert not any(x.id.startswith(('full_','opposite_')) for x in r.checks)

def test_interference_is_not_green_approval_and_still_exports():
    c=replace(next(iter(presets().values())),tool_radius=100.)
    r=evaluate(c)
    assert r.checks and all(x.passed for x in r.checks)
    assert r.status=='CÁLCULO ATENDE · CONFERIR INTERFERÊNCIAS'
    assert create_report(c,r)[:2]==b'PK'

def test_hole_outside_remaining_beam_section_still_blocks():
    c=replace(user_case(),plate_top=100.)
    r=evaluate(c)
    assert not r.checks and any('ligamento resistente' in i.text for i in r.issues)

def test_complement_word_contains_standard_checks_and_final_premise():
    c=replace(user_case(),plate_shape='between_flanges',root_width=45.)
    r=evaluate(c);d=Document(BytesIO(create_report(c,r)))
    text='\n'.join(p.text for p in d.paragraphs)
    cells='\n'.join(cell.text for t in d.tables for row in t.rows for cell in row.cells)
    assert text.index(COMPLEMENT_ASSUMPTION)>text.index('Premissas adotadas')
    assert 'Parafusos em corte simples' in cells and 'L calculado = 120,00' in cells
    assert 'SEM APROVAÇÃO DA LIGAÇÃO' not in text and 'Verificações parciais' not in text

def test_user_import_export_and_complement_toggle_in_streamlit():
    at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30).run()
    at.get('file_uploader')[0].upload('projeto.json',PROJECT.read_bytes(),'application/json').run()
    next(x for x in at.button if x.label=='Abrir arquivo').click().run()
    assert not at.exception and 'import_error' not in at.session_state
    assert at.selectbox(key='beam_name').value=='W 150 x 13,0'
    assert at.warning and not next(x for x in at.button if x.label=='Gerar memória Word').disabled
    at.selectbox(key='plate_shape').set_value('between_flanges').run()
    at.number_input(key='root_width').set_value(45.).run()
    next(x for x in at.button if x.label=='Gerar memória Word').click().run()
    assert not at.exception and at.session_state['report'][2][:2]==b'PK'
    at.selectbox(key='plate_shape').set_value('rectangular').run()
    at.selectbox(key='plate_shape').set_value('between_flanges').run()
    assert at.number_input(key='root_width').value==45.
    at.selectbox(key='kind').set_value('column_flange').run()
    assert not at.exception and at.selectbox(key='plate_shape').value=='rectangular'

def test_end_plate_access_interference_calculates_but_method_domain_blocks():
    from lro.end_plate import presets as ep_presets,evaluate as ep_evaluate
    c=replace(next(x for x in ep_presets().values() if x.kind=='moment'),tool_radius=45.)
    r=ep_evaluate(c)
    assert r.checks and any(i.severity=='interference' for i in r.issues)
    assert r.status=='CÁLCULO ATENDE · CONFERIR INTERFERÊNCIAS'
    assert not ep_evaluate(replace(c,N=1e8)).checks

def test_double_cover_tool_interference_preserves_calculation():
    from lro.double_cover import presets as dc_presets,evaluate as dc_evaluate
    c=next(iter(dc_presets().values()))
    base=dc_evaluate(c);r=dc_evaluate(replace(c,tool_radius=200.))
    assert r.checks==base.checks
    assert any(i.severity=='interference' for i in r.issues)
