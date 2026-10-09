"""Memória de end plate: figuras, dados, referências, contas e premissas finais."""
from io import BytesIO
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from .models import STEELS,KGF,VERSION
from .end_plate import NBR,DG,SCI
from .end_plate_drawing import image_bytes
from .report import table,number,display,decimal_text,hyperlink,BRAND,LINK,_math,group,sub,frac,power,radical

REFERENCES=[NBR+' — 6.1.2.3, 6.1.5.2, 6.2, 6.3, 6.5 e 5.7; coeficientes γa1=1,10 e γa2=γw2=1,35.',
            DG+' — 3.2, 3.5, 3.7.6, 5.1.1 e 5.3; tabelas 5-9 e 5-10. Modelo 4E de chapa espessa, sem nervuras.',
            SCI+' — recomendações de detalhamento para end plate rotulada de altura total; complemento de geometria, sem conversão de capacidades das tabelas.']

def equation(doc,row):
    i=row.id
    if i in ('ep_plate_rigidity','ep_column_rigidity'):
        _math(doc,group(sub('t','req'),' = ',radical(frac(group('1,10 ',sub('M','eq'),' max(',sub('γ','a1'),'; 1/0,90)'),group(sub('f','y'),' Y')))))
    elif i in ('ep_plate_prying','ep_column_prying'):
        _math(doc,group(sub('t','req'),' = ',radical(frac(group('4(b − ',sub('d','b'),'/2)',sub('F','t,0,Sd'),sub('γ','a1')),group('p ',sub('f','u'))))))
    elif i=='ep_bolt_interaction':
        _math(doc,group('η = ',power(group('(',frac(sub('F','t,Sd'),sub('F','t,Rd')),')')),' + ',power(group('(',frac(sub('F','v,Sd'),sub('F','v,Rd')),')')),' ≤ 1'))
    elif i=='ep_plate_yield':_math(doc,group(sub('M','Rd'),' = ',frac(group(sub('f','y'),power(sub('t','p')),sub('Y','p')),sub('γ','a1'))))
    else:_math(doc,group(row.equation))

def create_report(c,r,detailed=False):
    if not r.checks or any(x.severity=='error' for x in r.issues):raise ValueError('O projeto deve ter geometria e domínio válidos para gerar a memória.')
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2'):
        sty=doc.styles[name];sty.font.name='Arial';sty.font.color.rgb=RGBColor(0,0,0)
    for sty in doc.styles:
        for border in sty.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    doc.styles['Normal'].font.size=Pt(9.5);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(20);doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 2'].font.size=Pt(10)
    kind='rotulada' if c.kind=='pinned' else 'engastada'
    doc.core_properties.title='End plate '+kind+' na mesa de coluna';doc.core_properties.author='LRO Soluções de engenharia LTDA.'
    doc.add_paragraph('End plate '+kind,style='Title')
    doc.add_paragraph('Viga na mesa de coluna',style='Subtitle')
    doc.add_paragraph(c.project+' | '+c.beam.name+' → '+c.support.name)
    p=doc.add_paragraph();p.add_run(r.status).bold=True;p.add_run(' · maior índice '+number(max(x.ratio for x in r.checks),3))
    doc.add_picture(BytesIO(image_bytes(c)),width=Inches(6.95))
    doc.add_paragraph('A viga é soldada à chapa de topo, que se liga à mesa da coluna somente por parafusos. O desenho acompanha as entradas; as cotas são ilustrativas e devem ser complementadas no detalhamento de fabricação.')
    doc.add_paragraph('Ações no plano e já majoradas: V='+number(c.V/KGF)+' kgf; N='+number(c.N/KGF)+' kgf'+('; M='+number(c.M/(KGF*1000))+' kgf·m.' if c.kind=='moment' else '. Sem momento externo.'))
    doc.add_paragraph('N positivo indica tração; negativo, compressão. '+('M positivo traciona a mesa superior; negativo, a inferior.' if c.kind=='moment' else 'Os parafusos ficam entre as mesas, em grupo centrado.'))
    doc.add_page_break()
    doc.add_paragraph('Dimensões e materiais adotados',style='Heading 1')
    rows=[('Viga',f'd={c.beam.d:g}; bf={c.beam.bf:g}; tw={c.beam.tw:g}; tf={c.beam.tf:g}',c.beam_steel),
          ('Coluna',f'd={c.support.d:g}; bf={c.support.bf:g}; tw={c.support.tw:g}; tf={c.support.tf:g}',c.support_steel),
          ('Chapa',f'hp={c.hp:g}; bp={c.bp:g}; tp={c.tp:g}',c.plate_steel),
          ('Parafusos',f'{c.count} unidades; db={c.db:g}; g={c.gauge:g}',c.bolt),
          ('Furos',f'dh={c.dh:g}; dedução líquida={c.dn:g}','Padrão; '+('broca' if c.drilled else 'dedução +2 mm')),
          ('Linhas',', '.join(number(y,2) for y in c.rows),'y desde a face superior da viga'),
          ('Extensões',f'acima={c.ext_top:g}; abaixo={c.ext_bottom:g}','mm'),
          ('Soldas','CJP nas duas mesas e na alma' if c.kind=='moment' else f'Filetes contínuos w={c.weld:g}',f'fw={c.fw:g} MPa')]
    table(doc,['Componente','Dimensões em mm','Material ou critério'],[(a,decimal_text(b),d) for a,b,d in rows],[1.15,3.5,2.3])
    mats=list(dict.fromkeys([c.beam_steel,c.support_steel,c.plate_steel]))
    doc.add_paragraph('Resistências características: '+'; '.join(f'{s}: fy={STEELS[s].fy:g} MPa; fu={STEELS[s].fu:g} MPa' for s in mats)+'. E=200.000 MPa. 1 kgf=9,80665 N.')
    doc.add_paragraph('Referências de cálculo',style='Heading 1')
    for ref in REFERENCES:
        if (c.kind=='pinned' and ref.startswith(DG)) or (c.kind=='moment' and ref.startswith(SCI)):continue
        doc.add_paragraph(ref)
    doc.add_paragraph('Esforços usados nas verificações',style='Heading 1')
    a=r.geometry['actions'];aa=r.geometry['actual_actions']
    table(doc,['Caso','V em kgf','N em kgf','M em kgf·m'],[(lab,number(ac['V']/KGF),number(ac['N']/KGF),number(ac['M']/(1000*KGF))) for lab,ac in [('Entrada',aa),('Conferência governante',a)]],[2.3,1.55,1.55,1.55])
    doc.add_paragraph('A conferência mínima de 45 kN mantém a direção e o sentido da resultante N/V. Fator aplicado somente às forças: '+number(r.minimum_factor,4)+'. O momento externo não é majorado novamente. A entrada original também é verificada.')
    if c.kind=='moment':
        doc.add_paragraph('Distribuição local: Meq=|M|+max(N;0)(d−tf)/2; Ft por parafuso=Meq/[2(h1+h2)]. A compressão axial não reduz a tração de cálculo dos parafusos. O envelope é conservador; não representa uma análise elástica de contato.')
    doc.add_page_break()
    doc.add_paragraph('Quadro de verificações',style='Heading 1')
    doc.add_paragraph('Índice η=solicitação/resistência, ou interação; limite 1,000. Os critérios de espessura, ductilidade e compacidade são condições do método. O resultado abrange as verificações abaixo, sob as premissas finais.')
    vals=[]
    for row in r.checks:
        sd,unit=display(row.demand,row.unit);rd,_=display(row.resistance,row.unit)
        vals.append((row.name,sd,rd,unit,number(row.ratio,3),'Sim' if row.passed else 'Não'))
    table(doc,['Verificação','Sd ou valor','Rd ou limite','Unid.','η','Atende'],vals,[3.15,.9,.9,.6,.65,.75])
    doc.add_paragraph('Chapa: tp mínimo calculado='+number(a['tp_required'],3)+' mm; adotado='+number(c.tp,3)+' mm. Mesa da coluna: tf mínimo calculado='+number(a['column_tf_required'],3)+' mm; adotado='+number(c.support.tf,3)+' mm.')
    doc.add_page_break()
    doc.add_paragraph('Memória de cálculo',style='Heading 1')
    doc.add_paragraph('Substituições em N, mm e MPa. '+('Desenvolvimento de todas as verificações.' if detailed else 'Desenvolvimento dos critérios principais e dos itens governantes por componente. Todas as verificações constam do quadro.'))
    selected=r.checks if detailed else []
    if not detailed:
        groups=[['ep_plate_rigidity','ep_plate_prying'],['ep_column_rigidity','ep_column_prying'],['ep_rotation'],['ep_bolt_interaction'],['ep_bearing_plate','ep_bearing_column','ep_block_plate'],['ep_weld_fillet','ep_weld_cjp','ep_base_weld'],['ep_beam_web','ep_beam_nv'],['ep_column_web_yield','ep_column_crippling','ep_panel','ep_contact']]
        for ids in groups:
            candidates=[x for x in r.checks if x.id in ids]
            if candidates:selected.append(max(candidates,key=lambda x:x.ratio))
    for row in selected:
        start=len(doc.paragraphs);doc.add_paragraph(row.name,style='Heading 2');equation(doc,row)
        doc.add_paragraph(decimal_text(row.substitution)+'; η='+number(row.ratio,3)+'.')
        doc.add_paragraph(row.variables+' Referência: '+row.reference+'.')
        for p in doc.paragraphs[start:-1]:p.paragraph_format.keep_with_next=True
    if c.kind=='moment':
        doc.add_page_break()
        doc.add_paragraph('Parâmetros das linhas de plastificação',style='Heading 2')
        doc.add_paragraph('s=√(bp g)/2; pi=min(pfi;s); h1=d−tf/2+pfo; h2=d−1,5tf−pfi.')
        _math(doc,group(sub('Y','p'),' = (bp/2)[h1/pfo + h2(1/pi + 1/s) − 1/2] + (2/g)h2(s + pi)'))
        doc.add_paragraph('Para a coluna: sc=√(bfc g)/2; c=pfo+tf+pfi.')
        _math(doc,group(sub('Y','c'),' = (bfc/2)(h1+h2)/sc + (2/g)[h2(3c/4+sc)+h1(c/4+sc)+c²/2] + g/2'))
        doc.add_paragraph('Valores: h1='+number(a['h1'],3)+'; h2='+number(a['h2'],3)+'; Yp='+number(a['Yp'],3)+' mm; Yc='+number(a['Yc'],3)+' mm. A espessura mínima controla a hipótese de chapa espessa e alavanca desprezível; não é, isoladamente, uma medição de rigidez global.')
    doc.add_paragraph('Variáveis',style='Heading 2')
    doc.add_paragraph('d, bf, tw e tf: altura, largura de mesa, alma e mesa da viga; bp, hp e tp: largura, altura e espessura da chapa; g: gabarito horizontal; pfi e pfo: distâncias livres da mesa às linhas de parafusos; db e dh: diâmetros do parafuso e furo; Ab: área nominal do parafuso; fy, fu e fub: resistências características; Ft e Fv: forças de tração e cisalhamento por parafuso; lc: ligamento livre; Cv: fator de cisalhamento da alma; Sd e Rd: solicitação e resistência de cálculo. Índice c identifica a coluna; p identifica a chapa.')
    if c.notes:doc.add_paragraph('Observações do projeto: '+c.notes)
    doc.add_paragraph('Premissas adotadas',style='Heading 1')
    for note in c.assumptions:doc.add_paragraph(note)
    for issue in r.issues:doc.add_paragraph(issue.text)
    doc.add_paragraph(BRAND)
    hyperlink(doc.add_paragraph('LRO Ligações '+VERSION+' · '),'LinkedIn de Lucas Oliveira',LINK)
    out=BytesIO();doc.save(out);return out.getvalue()
